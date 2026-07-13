"""CSE build-flow schematic — faithful to the original 7-stage left-to-right
layout, redrawn with aligned stages, complete arrows, and lane-kind colors."""
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x0F, 0x17, 0x2A)
TEXT = RGBColor(0x33, 0x41, 0x55)
MUTED = RGBColor(0x64, 0x74, 0x8B)
PANEL = RGBColor(0xF1, 0xF5, 0xF9)
BORDER = RGBColor(0xCB, 0xD5, 0xE1)
NAVY = RGBColor(0x1E, 0x3A, 0x8A)      # titles / stage circles / build flow
TEAL = RGBColor(0x0D, 0x94, 0x88)      # core + reuse edges
AMBER = RGBColor(0xD9, 0x77, 0x06)     # serial
BLUE = RGBColor(0x25, 0x63, 0xEB)      # mpi
VIOLET = RGBColor(0x7C, 0x3A, 0xED)    # gpu
GREEN = RGBColor(0x15, 0x80, 0x3D)     # gcc badge
ORANGE = RGBColor(0xC2, 0x41, 0x0C)    # second compiler badge
GRAY = RGBColor(0x47, 0x55, 0x69)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]

TINT = {AMBER: RGBColor(0xFD, 0xF3, 0xE3), BLUE: RGBColor(0xEA, 0xF0, 0xFE),
        VIOLET: RGBColor(0xF2, 0xEB, 0xFD), GRAY: RGBColor(0xF1, 0xF5, 0xF9)}


def txt(s, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    b = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = b.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    first = True
    for sz, col, bd, t in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.text = t; p.alignment = align
        p.font.size = Pt(sz); p.font.color.rgb = col; p.font.bold = bd
        p.font.name = "Avenir Next"
    return b


def box(s, x, y, w, h, fill, line=None, lw=1.25, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    shp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    shp.shadow.inherit = False
    if line:
        shp.line.color.rgb = line; shp.line.width = Pt(lw)
    else:
        shp.line.fill.background()
    return shp


def conn(s, x1, y1, x2, y2, color, w=1.75, head=True, dashed=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                               Inches(x2), Inches(y2))
    c.line.color.rgb = color; c.line.width = Pt(w)
    c.shadow.inherit = False
    ln = c.line._get_or_add_ln()
    if dashed:
        d = etree.SubElement(ln, qn('a:prstDash')); d.set('val', 'dash')
    if head:
        t = etree.SubElement(ln, qn('a:tailEnd'))
        t.set('type', 'triangle'); t.set('w', 'med'); t.set('len', 'med')
    return c


def stage(s, n, label, cx, y=0.98):
    c = box(s, cx - 0.14, y, 0.28, 0.28, NAVY, shape=MSO_SHAPE.OVAL)
    tf = c.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.text = str(n); p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(11); p.font.bold = True; p.font.color.rgb = WHITE
    p.font.name = "Avenir Next"
    txt(s, cx - 0.85, y + 0.31, 1.7, 0.42, [(8.5, INK, True, label)],
        align=PP_ALIGN.CENTER)


def pentagon(s, x, y, name, sub, color):
    box(s, x + 0.42, y, 0.5, 0.42, color, shape=MSO_SHAPE.REGULAR_PENTAGON)
    txt(s, x, y + 0.46, 1.34, 0.55, [(13, color, True, name), (8, MUTED, False, sub)],
        align=PP_ALIGN.CENTER)


def build_slide(title, subtitle, facts, compilers, cards, modules, callout_naming,
                callout_mpi):
    s = prs.slides.add_slide(BLANK)
    box(s, 0, 0, 13.333, 7.5, WHITE, shape=MSO_SHAPE.RECTANGLE)
    txt(s, 0.4, 0.12, 12.5, 0.5, [(23, INK, True, title)], align=PP_ALIGN.CENTER)
    txt(s, 0.4, 0.6, 12.5, 0.35, [(12.5, BLUE, True, subtitle)], align=PP_ALIGN.CENTER)

    centers = [1.2, 3.05, 4.72, 6.55, 9.55, 11.45, 12.66]
    labels = ["System inspection", "Register compilers", "Compilers ready",
              "Build Core once", "Fan-out lanes", "Build and validate",
              "Publish to users"]
    for i, (cx, lb) in enumerate(zip(centers, labels)):
        stage(s, i + 1, lb, cx)
        if i < 6:
            conn(s, cx + 0.55, 1.12, centers[i + 1] - 0.55, 1.12, MUTED, w=1.25)

    top, bot = 1.85, 6.55

    # 1 · system facts
    box(s, 0.32, top, 1.76, bot - top, PANEL, line=BORDER)
    txt(s, 0.44, top + 0.1, 1.55, 0.55, [(9.5, NAVY, True, "🔍  Environment facts (probed)")])
    for i, f in enumerate(facts):
        b = box(s, 0.44, top + 0.72 + i * 0.56, 1.52, 0.46, WHITE, line=BORDER)
        tf = b.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = Inches(0.05)
        p = tf.paragraphs[0]; p.text = f; p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(8.5); p.font.color.rgb = TEXT; p.font.name = "Avenir Next"

    # 2 · compilers
    box(s, 2.22, top, 1.66, bot - top, PANEL, line=BORDER)
    txt(s, 2.34, top + 0.1, 1.45, 0.55, [(9.5, NAVY, True, "Compiler surfaces (externals)")])
    pentagon(s, 2.38, top + 0.85, *compilers[0])
    pentagon(s, 2.38, top + 2.35, *compilers[1])
    txt(s, 2.34, bot - 0.75, 1.45, 0.65,
        [(8.5, MUTED, False, "each surface = its own module tree")],
        align=PP_ALIGN.CENTER)

    # 3 · compilers ready
    box(s, 4.12, 3.3, 1.2, 1.5, WHITE, line=VIOLET)
    txt(s, 4.18, 3.42, 1.08, 1.3, [
        (16, INK, False, "🛠"),
        (8.5, VIOLET, True, "compilers registered ·"),
        (8.5, VIOLET, True, "toolchains defined"),
    ], align=PP_ALIGN.CENTER)

    # 4 · core environment
    box(s, 5.55, top, 2.0, bot - top, WHITE, line=NAVY, lw=1.5)
    txt(s, 5.7, top + 0.1, 1.72, 4.6, [
        (9.5, NAVY, True, "📦  Spack Core environment"),
        (8.5, TEXT, False, ""),
        (8.5, TEXT, False, "• built with GCC"),
        (8.5, TEXT, False, "• portable target: x86_64_v3"),
        (8.5, TEXT, False, "• cmake · ninja · git"),
        (8.5, TEXT, False, "• python · miniforge · gsl"),
        (8.5, TEXT, False, ""),
        (8.5, GRAY, True, "foundation — in the view:"),
        (8.5, TEXT, False, "• zlib · xz · zstd (pinned)"),
        (8.5, TEXT, False, ""),
        (9, TEAL, True, "shared across all surfaces and lanes"),
    ])

    # panel-to-panel flow arrows (build flow, navy; core→cache is reuse, teal)
    conn(s, 2.08, 4.2, 2.22, 4.2, NAVY, w=1.25)
    conn(s, 3.88, 4.05, 4.12, 4.05, NAVY, w=1.25)
    conn(s, 5.32, 4.05, 5.55, 4.05, NAVY, w=1.25)
    conn(s, 7.55, 4.05, 7.72, 4.05, TEAL, w=1.25)

    # cache node
    box(s, 7.72, 3.62, 0.62, 0.86, TINT[GRAY], line=TEAL)
    txt(s, 7.71, 3.7, 0.66, 0.72, [(11, INK, False, "🗂"), (7, TEAL, True, "core view + cache")],
        align=PP_ALIGN.CENTER)

    # 5 · lane cards
    n = len(cards)
    ch = min(0.72, (bot - top - 0.08 * (n - 1)) / n)
    spine_x = 8.52
    for i, (name, content, color) in enumerate(cards):
        y = top + i * (ch + 0.08)
        box(s, 8.62, y, 2.42, ch, TINT[color], line=color)
        txt(s, 8.74, y + 0.06, 1.85, ch - 0.1, [
            (9.5, color, True, name),
            (7.5, TEXT, False, content),
        ])
        txt(s, 10.6, y + 0.06, 0.42, ch - 0.1,
            [(6.5, MUTED, False, "yaml lock view mods")])
        conn(s, spine_x, y + ch / 2, 8.62, y + ch / 2, TEAL, w=1.5)
        conn(s, 11.04, y + ch / 2, 11.18, y + ch / 2, NAVY, w=1.25, head=False)
    # reuse spine from cache to cards
    conn(s, 8.34, 4.05, spine_x, 4.05, TEAL, w=1.75, head=False, dashed=True)
    conn(s, spine_x, top + ch / 2, spine_x, top + (n - 1) * (ch + 0.08) + ch / 2,
         TEAL, w=1.75, head=False, dashed=True)
    # build-flow collector into validate
    conn(s, 11.18, top + ch / 2, 11.18, top + (n - 1) * (ch + 0.08) + ch / 2,
         NAVY, w=1.25, head=False)
    conn(s, 11.18, 4.05, 11.32, 4.05, NAVY, w=1.75)

    # 6 · build and validate
    box(s, 11.32, top, 0.95, bot - top, PANEL, line=TEAL)
    steps = ["Concretize", "Fetch", "Install", "Test / verify", "Benchmarks", "Smoke runs"]
    txt(s, 11.4, top + 0.1, 0.8, 0.35, [(11, TEAL, False, "✅")], align=PP_ALIGN.CENTER)
    for i, st in enumerate(steps):
        txt(s, 11.38, top + 0.5 + i * 0.62, 0.84, 0.4, [(8.5, TEXT, True, st)],
            align=PP_ALIGN.CENTER)
        if i < 5:
            conn(s, 11.55, top + 0.92 + i * 0.62, 11.55, top + 1.06 + i * 0.62,
                 TEAL, w=1.0)
    txt(s, 11.38, bot - 0.45, 0.84, 0.42,
        [(7, MUTED, False, "lanes build in parallel")], align=PP_ALIGN.CENTER)
    conn(s, 12.27, 4.05, 12.4, 4.05, NAVY, w=1.75)

    # 7 · publish
    box(s, 12.4, top, 0.6, bot - top, PANEL, line=VIOLET)
    txt(s, 12.42, top + 0.1, 0.56, 0.4, [(12, INK, False, "📣")], align=PP_ALIGN.CENTER)
    for i, t in enumerate(["views", "modules", "caches", "release"]):
        txt(s, 12.4, top + 0.62 + i * 0.5, 0.6, 0.35, [(8, VIOLET, True, t)],
            align=PP_ALIGN.CENTER)
    # user modules list under publish, spanning bottom-right
    txt(s, 11.32, bot + 0.06, 1.7, 0.85,
        [(7.5, INK, True, "user modules")] + [(7, MUTED, False, m) for m in modules[:4]])

    # legend + callouts
    txt(s, 0.34, bot + 0.14, 0.9, 0.3, [(8.5, INK, True, "legend")])
    conn(s, 1.15, bot + 0.26, 1.55, bot + 0.26, NAVY, w=1.75)
    txt(s, 1.6, bot + 0.14, 1.0, 0.3, [(8, MUTED, False, "build flow")])
    conn(s, 2.6, bot + 0.26, 3.0, bot + 0.26, TEAL, w=1.75, dashed=True, head=False)
    txt(s, 3.05, bot + 0.14, 1.3, 0.3, [(8, MUTED, False, "reuse — built once")])
    b = box(s, 4.7, bot + 0.08, 3.6, 0.5, TINT[BLUE], line=BLUE)
    tf = b.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.text = callout_naming; p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(8.5); p.font.bold = True; p.font.color.rgb = NAVY
    p.font.name = "Avenir Next"
    b = box(s, 8.5, bot + 0.08, 2.6, 0.5, TINT[VIOLET], line=VIOLET)
    tf = b.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.text = callout_mpi; p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(8.5); p.font.bold = True; p.font.color.rgb = VIOLET
    p.font.name = "Avenir Next"
    return s


build_slide(
    "CSE Spack Stack Build Flow — Cray System",
    "Cray PE compiler surfaces + shared Core + fan-out Cray MPICH lanes",
    ["OS = Cray Linux (SLES)", "CPU target = zen4", "Cray PE / PrgEnv",
     "MPI = Cray MPICH", "Fabric = Slingshot / CXI", "GPU = MI300A (gfx942)",
     "Scheduler = PBS / Slurm"],
    [("GCC", "PrgEnv-gnu · default", GREEN), ("CCE", "PrgEnv-cray · vendor", ORANGE)],
    [("GCC / Serial", "HDF5 · NetCDF — built without MPI", AMBER),
     ("GCC / MPI", "Cray MPICH · HDF5+MPI · PnetCDF · TAU", BLUE),
     ("GCC / GPU", "Kokkos · ROCm gfx942 · GPU-aware MPICH", VIOLET),
     ("CCE / Serial", "HDF5 · NetCDF — built without MPI", AMBER),
     ("CCE / MPI", "Cray MPICH · HDF5+MPI · PnetCDF · TAU", BLUE),
     ("CCE / GPU", "Kokkos · ROCm gfx942 · GPU-aware MPICH", VIOLET)],
    ["CSE/Core (auto)", "CSE/GCC/Serial", "CSE/GCC/MPI", "CSE/GCC/GPU"],
    "One MPI on the system → lanes stay unqualified: Serial · MPI · GPU",
    "Cray MPICH is the production MPI",
)

build_slide(
    "CSE Spack Stack Build Flow — Linux System",
    "Site compiler surfaces + shared Core + lanes qualified per MPI implementation",
    ["OS = RHEL 9", "CPU target = zen3", "Site MPIs = Open MPI + MPICH",
     "Fabric = InfiniBand / UCX", "GPU = A100 (sm_80)", "Scheduler = Slurm",
     "Modules = Lmod"],
    [("GCC", "site default", GREEN), ("AOCC", "external · optional", ORANGE)],
    [("GCC / Serial", "HDF5 · NetCDF — built without MPI", AMBER),
     ("GCC / MPI-openmpi", "Open MPI 5 (UCX) · HDF5+MPI · TAU", BLUE),
     ("GCC / MPI-mpich", "MPICH 4 · HDF5+MPI · TAU", BLUE),
     ("GCC / GPU-openmpi", "Kokkos · CUDA sm_80 · over Open MPI", VIOLET),
     ("GCC / GPU-mpich", "Kokkos · CUDA sm_80 · over MPICH", VIOLET),
     ("AOCC / …", "same lane set per compiler surface", GRAY)],
    ["CSE/Core (auto)", "CSE/GCC/Serial", "CSE/GCC/MPI-openmpi", "CSE/GCC/GPU-openmpi"],
    "Two MPIs on the system → lanes qualify: MPI-<impl> · GPU-<impl>",
    "Open MPI typically built with UCX",
)

out = Path(__file__).with_name("cse_build_flow.pptx")
prs.save(out)
print("WROTE", out)
